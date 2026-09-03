import json
from pathlib import Path

from docx import Document

root = Path(__file__).resolve().parents[2]
calendar_path = root / "南京工业大学2026-2027学年校历.xls"
template_path = root / "周一汇报-模版.docx"

doc = Document(template_path)
template = {
    "paragraphs": [p.text.strip() for p in doc.paragraphs if p.text.strip()],
    "tables": [
        [[cell.text.strip() for cell in row.cells] for row in table.rows]
        for table in doc.tables
    ],
}

print(json.dumps({"template": template}, ensure_ascii=False, indent=2))
